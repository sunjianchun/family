#!/usr/bin/python
# -*- coding: utf-8 -*-
from docx.api import Document
import re, json
doc = Document("/Users/sunjianchun/Downloads/家谱四稿.docx")
tables = doc.tables

def ReadWord():
	newtables = []
	dict = {}
	name_pattern = re.compile(r'姓[ \t\n]*名[ \t\n]*$')
	fellowRank_pattern = re.compile(r'同[ \t\n]*辈[ \t\n]*总[ \t\n]*行')
	compatriotRank_pattern = re.compile(r'同[ \t\n]*胞[ \t\n]*行')
	selfIntroduce_pattern = re.compile(r'本[ \t\n]*人[ \t\n]*简[ \t\n]*历')
	age_pattern = re.compile(r'年[ \t\n]*龄')
	birthday_pattern = re.compile(r'出[ \t\n]*生[ \t\n]*年[ \t\n]*月')
	spouseIntroduce_pattern = re.compile(r'配[ \t\n]*偶[ \t\n]*简[ \t\n]*历')
	parents_pattern = re.compile(r'父[ \t\n]*母')
	brother_pattern = re.compile(r'同[ \t\n]*胞[ \t\n]*兄[ \t\n]*弟')
	sisters_pattern = re.compile(r'同[ \t\n]*胞[ \t\n]*姐[ \t\n]*妹')
	children_pattern = re.compile(r'子[ \t\n]*女')
	remark_pattern = re.compile(r'备[ \t\n]*注')
	phone_pattern = re.compile(r'电[ \t\n]*话')
	
	dict["name_pattern"] = name_pattern
	dict["fellowRank_pattern"] = fellowRank_pattern
	dict["compatriotRank_pattern"] = compatriotRank_pattern
	dict["age_pattern"] = age_pattern
	dict["selfIntroduce_pattern"] = selfIntroduce_pattern
	dict["birthday_pattern"] = birthday_pattern
	dict["spouseIntroduce_pattern"] = spouseIntroduce_pattern
	dict["parents_pattern"] = parents_pattern
	dict["brother_pattern"] = brother_pattern
	dict["sisters_pattern"] = sisters_pattern
	dict["children_pattern"] = children_pattern
	dict["remark_pattern"] = remark_pattern
	dict["phone_pattern"] = phone_pattern
	#tables = tables[202:203]
	
	for t in tables:
	    for i in range(len(t.rows)):
	        for j in range(len(t.columns)):
	            try:
	                if selfIntroduce_pattern.match((t.cell(i, j).text).encode('utf-8')):
	                    newtables.append(t)
	                    break
	                else:
	                    continue
	            except Exception, e:
	                continue
	newtables = list(set(newtables))
	
	#newtables = newtables[20:30]
	newlist = []
	f = open('/tmp/file.txt', 'wb')
	for t in newtables:
	    newdict = {}
	    for i in range(len(t.rows)):
	        for j in range(len(t.columns)):
	            try:
	                for k, v in dict.items():
	                    if v.match((t.cell(i, j).text).encode('utf-8')) and v.match((t.cell(i, j+1).text).encode('utf-8')):
	                        continue
	                    elif v.match((t.cell(i, j).text).encode('utf-8')) and not v.match((t.cell(i, j+1).text).encode('utf-8')):
	                        newdict[k.split("_")[0]] = (t.cell(i, j+1).text)
	                        break
	                    else:
	                        continue
	            except Exception, e:
	                continue
	            j += 1           
	    newlist.append(newdict)
	a = json.dumps(newlist, sort_keys=True, indent=4, separators=(',', ': ')).encode('utf-8')
	f.write(a)
	f.close()
if __name__ == '__main__':
	ReadWord()
